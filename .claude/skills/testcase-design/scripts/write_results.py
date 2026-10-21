"""Ghi kết quả chạy thử (pretest) vào file Excel test case đã tạo.

    .venv/bin/python -I .claude/skills/testcase-design/scripts/write_results.py \
        --xlsx output/TestCases_<ten>_<YYYYMMDD>.xlsx --results _workspace/05_pretest_results.json

results.json: {"run_label": "Lần 1", "executed_by": "Claude in Chrome",
               "results": [{"tc_id","status","actual","note","executed_at"?}]}
Chỉ cập nhật các cột Trạng thái / Kết quả thực tế / Thời điểm / Người chạy / Ghi chú của
đúng dòng TC ID; không đổi nội dung thiết kế test case. Thoát mã 3 nếu có chuỗi nhạy cảm.
"""
import argparse
import json
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xlsx_common import STATUSES, fail_on_secrets  # noqa: E402

from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import Alignment  # noqa: E402

HEADERS = {"tc_id": "TC ID", "status": "Trạng thái", "actual": "Kết quả thực tế",
           "executed_at": "Thời điểm chạy", "executed_by": "Người/Công cụ chạy", "notes": "Ghi chú / Lỗi"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--results", required=True)
    a = ap.parse_args()
    with open(a.results, encoding="utf-8") as f:
        payload = json.load(f)
    fail_on_secrets(payload, os.path.basename(a.results))

    bad = [r for r in payload.get("results", []) if r.get("status") not in STATUSES or r.get("status") == "Not Run"]
    if bad:
        print("[LỖI] status không hợp lệ (chỉ Pass/Fail/Blocked/Skipped):", [b.get("tc_id") for b in bad], file=sys.stderr)
        sys.exit(2)

    wb = load_workbook(a.xlsx)
    ws = wb["Test Cases"]
    header = {c.value: c.column for c in ws[1]}
    cols = {k: header[v] for k, v in HEADERS.items()}
    rows = {ws.cell(r, cols["tc_id"]).value: r for r in range(2, ws.max_row + 1)}
    log = wb["Pretest Log"]
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    label = payload.get("run_label", "Lần chạy")
    by = payload.get("executed_by", "Claude in Chrome")
    updated, missing = 0, []
    for r in payload.get("results", []):
        row = rows.get(r["tc_id"])
        if not row:
            missing.append(r["tc_id"])
            continue
        ts = r.get("executed_at") or now
        ws.cell(row, cols["status"]).value = r["status"]
        ws.cell(row, cols["actual"]).value = r.get("actual", "")
        ws.cell(row, cols["executed_at"]).value = ts
        ws.cell(row, cols["executed_by"]).value = by
        if r.get("note"):
            ws.cell(row, cols["notes"]).value = r["note"]
        for k in ("status", "actual", "executed_at", "executed_by", "notes"):
            ws.cell(row, cols[k]).alignment = Alignment(wrap_text=True, vertical="top")
        log.append([label, ts, r["tc_id"], r["status"], r.get("note") or r.get("actual", "")])
        updated += 1
    wb.save(a.xlsx)
    print(f"Đã cập nhật {updated} test case vào {a.xlsx}")
    if missing:
        print("Không tìm thấy TC ID:", ", ".join(missing), file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()
