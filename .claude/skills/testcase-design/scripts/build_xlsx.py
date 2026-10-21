"""Tạo file Excel test case từ _workspace/03_testcases.json.

Cách chạy (từ thư mục gốc dự án):
    .venv/bin/python -I .claude/skills/testcase-design/scripts/build_xlsx.py \
        --input _workspace/03_testcases.json --output output/TestCases_<ten>_<YYYYMMDD>.xlsx

Các sheet: Tổng quan, Test Cases, Traceability, Spec Mapping, Test Data, Pretest Log.
Thoát với mã 2 nếu dữ liệu sai cấu trúc, mã 3 nếu phát hiện chuỗi nhạy cảm.
"""
import argparse
import json
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xlsx_common import STATUSES, PRIORITIES, MAPPING_STATUSES, fail_on_secrets  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

TC_COLUMNS = [
    ("tc_id", "TC ID", 12),
    ("module", "Module", 16),
    ("feature", "Chức năng", 20),
    ("req_ids", "Req ID", 14),
    ("spec_ref", "Tham chiếu spec", 26),
    ("title", "Tiêu đề", 38),
    ("type", "Loại", 14),
    ("priority", "Ưu tiên", 9),
    ("technique", "Kỹ thuật thiết kế", 18),
    ("preconditions", "Điều kiện tiên quyết", 30),
    ("test_data", "Dữ liệu test", 28),
    ("steps", "Các bước thực hiện", 48),
    ("expected", "Kết quả mong đợi", 42),
    ("chrome_testable", "Chạy được bằng Chrome", 12),
    ("status", "Trạng thái", 12),
    ("actual", "Kết quả thực tế", 36),
    ("executed_at", "Thời điểm chạy", 18),
    ("executed_by", "Người/Công cụ chạy", 16),
    ("notes", "Ghi chú / Lỗi", 32),
]
REQUIRED_TC = ["tc_id", "module", "title", "type", "priority", "steps", "expected", "req_ids"]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def as_text(v):
    if v is None:
        return ""
    if isinstance(v, list):
        if all(isinstance(x, str) for x in v):
            return "\n".join(f"{i}. {x}" if len(v) > 1 else x for i, x in enumerate(v, 1))
        return "\n".join(str(x) for x in v)
    if isinstance(v, bool):
        return "Y" if v else "N"
    return str(v)


def style_header(ws, row=1):
    for c in ws[row]:
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 32


def validate(data):
    errs = []
    tcs = data.get("test_cases") or []
    if not tcs:
        errs.append("test_cases rỗng")
    seen = set()
    req_ids = {r.get("id") for r in data.get("requirements", [])}
    for i, tc in enumerate(tcs):
        for k in REQUIRED_TC:
            if not tc.get(k):
                errs.append(f"test_cases[{i}] ({tc.get('tc_id','?')}) thiếu '{k}'")
        tid = tc.get("tc_id")
        if tid in seen:
            errs.append(f"TC ID trùng: {tid}")
        seen.add(tid)
        if tc.get("priority") not in PRIORITIES:
            errs.append(f"{tid}: priority '{tc.get('priority')}' không thuộc {PRIORITIES}")
        if tc.get("status", "Not Run") not in STATUSES:
            errs.append(f"{tid}: status không hợp lệ")
        for rid in (tc.get("req_ids") or []) if isinstance(tc.get("req_ids"), list) else []:
            if req_ids and rid not in req_ids:
                errs.append(f"{tid}: req_ids '{rid}' không có trong requirements")
    for r in data.get("requirements", []):
        if r.get("mapping_status") and r["mapping_status"] not in MAPPING_STATUSES:
            errs.append(f"{r.get('id')}: mapping_status '{r['mapping_status']}' không hợp lệ")
    covered = {rid for tc in tcs for rid in (tc.get("req_ids") if isinstance(tc.get("req_ids"), list) else [])}
    for r in data.get("requirements", []):
        if r.get("id") not in covered and r.get("mapping_status") != "SPEC_ONLY":
            errs.append(f"Requirement {r.get('id')} chưa có test case nào phủ")
    return errs


def build(data, out_path):
    wb = Workbook()
    meta = data.get("meta", {})

    # --- Test Cases ---
    ws = wb.active
    ws.title = "Test Cases"
    ws.append([c[1] for c in TC_COLUMNS])
    style_header(ws)
    for tc in data["test_cases"]:
        row = []
        for key, _, _ in TC_COLUMNS:
            v = tc.get(key)
            if key == "status" and not v:
                v = "Not Run"
            if key == "chrome_testable" and v is None:
                v = "N"
            if key == "req_ids" and isinstance(v, list):
                v = ", ".join(v)
            row.append(as_text(v))
        ws.append(row)
    for idx, (_, _, w) in enumerate(TC_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment, c.border = WRAP, BORDER
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = ws.dimensions
    n = len(data["test_cases"]) + 1
    col = {k: get_column_letter(i) for i, (k, _, _) in enumerate(TC_COLUMNS, 1)}
    dv_s = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    dv_p = DataValidation(type="list", formula1='"' + ",".join(PRIORITIES) + '"', allow_blank=False)
    dv_c = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    for dv, k in ((dv_s, "status"), (dv_p, "priority"), (dv_c, "chrome_testable")):
        ws.add_data_validation(dv)
        dv.add(f"{col[k]}2:{col[k]}{max(n, 2) + 500}")
    rng = f"{col['status']}2:{col['status']}{max(n, 2) + 500}"
    for val, color in (("Pass", "C6EFCE"), ("Fail", "FFC7CE"), ("Blocked", "FFEB9C"), ("Skipped", "D9D9D9")):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", bgColor=color)))

    # --- Tổng quan ---
    ov = wb.create_sheet("Tổng quan", 0)
    ov.column_dimensions["A"].width = 30
    ov.column_dimensions["B"].width = 60
    info = [
        ("Tên bộ test", meta.get("title", "")),
        ("Yêu cầu đầu vào", meta.get("input_summary", "")),
        ("Tài liệu spec đã đối chiếu", ", ".join(meta.get("spec_files", []))),
        ("Ngày tạo", meta.get("created") or datetime.now().strftime("%Y-%m-%d")),
        ("Môi trường test (nếu có)", meta.get("env_note", "Chưa xác định")),
        ("Giả định / điểm đã được người dùng xác nhận", meta.get("assumptions", "")),
    ]
    for k, v in info:
        ov.append([k, v])
        ov.cell(ov.max_row, 1).font = Font(bold=True)
        ov.cell(ov.max_row, 2).alignment = WRAP
    ov.append([])
    ov.append(["Thống kê", "Số lượng"])
    style_header(ov, ov.max_row)
    last = max(n, 2) + 500
    st = f"'Test Cases'!${col['status']}$2:${col['status']}${last}"
    pr = f"'Test Cases'!${col['priority']}$2:${col['priority']}${last}"
    idc = f"'Test Cases'!$A$2:$A${last}"
    ov.append(["Tổng số test case", f'=COUNTA({idc})'])
    for s in STATUSES:
        ov.append([f"Trạng thái: {s}", f'=COUNTIF({st},"{s}")'])
    for p in PRIORITIES:
        ov.append([f"Ưu tiên: {p}", f'=COUNTIF({pr},"{p}")'])
    ov.append(["Tỉ lệ Pass (trên số đã chạy)", f'=IFERROR(COUNTIF({st},"Pass")/(COUNTIF({st},"Pass")+COUNTIF({st},"Fail")+COUNTIF({st},"Blocked")),"-")'])
    ov.cell(ov.max_row, 2).number_format = "0.0%"

    # --- Traceability ---
    tr = wb.create_sheet("Traceability")
    tr.append(["Req ID", "Nội dung yêu cầu", "Tham chiếu spec", "Kết quả đối chiếu", "Số TC", "Danh sách TC"])
    style_header(tr)
    for r in data.get("requirements", []):
        tcs = [tc["tc_id"] for tc in data["test_cases"] if r["id"] in (tc.get("req_ids") if isinstance(tc.get("req_ids"), list) else [])]
        tr.append([r.get("id"), r.get("text"), r.get("spec_ref", ""), r.get("mapping_status", ""), len(tcs), ", ".join(tcs)])
    for i, w in enumerate((10, 60, 30, 18, 8, 50), 1):
        tr.column_dimensions[get_column_letter(i)].width = w
    for row in tr.iter_rows(min_row=2):
        for c in row:
            c.alignment, c.border = WRAP, BORDER
    tr.freeze_panes = "A2"

    # --- Spec Mapping ---
    sm = wb.create_sheet("Spec Mapping")
    sm.append(["Req ID", "Kết quả đối chiếu", "Điểm lệch / cần xác nhận", "Trích dẫn spec", "Câu hỏi đã hỏi", "Quyết định của người dùng"])
    style_header(sm)
    for r in data.get("requirements", []):
        if r.get("mapping_status") in (None, "MATCH"):
            continue
        sm.append([r.get("id"), r.get("mapping_status"), r.get("discrepancy", ""), r.get("spec_quote", ""), r.get("question", ""), r.get("resolution", "(chưa có)")])
    for i, w in enumerate((10, 18, 50, 50, 40, 40), 1):
        sm.column_dimensions[get_column_letter(i)].width = w
    for row in sm.iter_rows(min_row=2):
        for c in row:
            c.alignment, c.border = WRAP, BORDER
    sm.freeze_panes = "A2"

    # --- Test Data ---
    td = wb.create_sheet("Test Data")
    td.append(["Tên dữ liệu", "Mô tả", "Giá trị / quy ước (chỉ placeholder, KHÔNG dữ liệu thật)"])
    style_header(td)
    for d in data.get("test_data", []):
        td.append([d.get("name"), d.get("description"), d.get("value")])
    for i, w in enumerate((28, 50, 60), 1):
        td.column_dimensions[get_column_letter(i)].width = w
    for row in td.iter_rows(min_row=2):
        for c in row:
            c.alignment, c.border = WRAP, BORDER

    # --- Pretest Log ---
    pl = wb.create_sheet("Pretest Log")
    pl.append(["Lần chạy", "Thời điểm", "TC ID", "Kết quả", "Ghi chú"])
    style_header(pl)
    for i, w in enumerate((10, 20, 12, 12, 80), 1):
        pl.column_dimensions[get_column_letter(i)].width = w

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--output", required=True)
    a = ap.parse_args()
    with open(a.input, encoding="utf-8") as f:
        data = json.load(f)
    errs = validate(data)
    if errs:
        print("[LỖI CẤU TRÚC]", file=sys.stderr)
        for e in errs[:50]:
            print("  -", e, file=sys.stderr)
        sys.exit(2)
    fail_on_secrets(data, os.path.basename(a.input))
    build(data, a.output)
    print(f"Đã tạo {a.output}: {len(data['test_cases'])} test case, {len(data.get('requirements', []))} requirement")


if __name__ == "__main__":
    main()
